import docx
import os
from docx.shared import Inches

def read_docx_content(file_path):
    try:
        doc = docx.Document(file_path)
        content = []
        
        print(f"\n--- {file_path} 内容 ---\n")
        
        # 读取段落
        for i, para in enumerate(doc.paragraphs):
            if para.text.strip():
                print(f"段落 {i+1}: {para.text}")
                content.append(para.text)
        
        # 读取表格
        print("\n--- 表格内容 ---\n")
        for i, table in enumerate(doc.tables):
            print(f"表格 {i+1}:")
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                print(f"  行: {row_text}")
                content.extend(row_text)
        
        # 读取图片（仅获取数量，无法直接提取图片内容）
        print("\n--- 图片信息 ---\n")
        images = []
        for rel in doc.part.rels.values():
            if "image" in rel.target_ref:
                images.append(rel.target_ref)
        print(f"文档中包含 {len(images)} 张图片")
        
        return content
    except Exception as e:
        print(f"Error reading file: {str(e)}")
        return []

# 读取文件
file_path = '2.7改.docx'
content = read_docx_content(file_path)

print("\n--- 提取的内容 ---\n")
for line in content:
    print(line)
